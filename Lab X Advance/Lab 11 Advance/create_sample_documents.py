# create_sample_documents.py
"""Creates the sample input files: 3 PDFs in documents/ and the Word template in templates/."""
import os

import docx
from docx.shared import Pt, RGBColor
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

BASE = os.path.dirname(os.path.abspath(__file__))
DOCS = os.path.join(BASE, "documents")
TEMPLATES = os.path.join(BASE, "templates")
os.makedirs(DOCS, exist_ok=True)
os.makedirs(TEMPLATES, exist_ok=True)


def make_pdf(filename, title, lines):
    c = canvas.Canvas(os.path.join(DOCS, filename), pagesize=A4)
    w, h = A4
    c.setFillColorRGB(0.12, 0.31, 0.55)
    c.rect(0, h - 90, w, 90, fill=1, stroke=0)
    c.setFillColorRGB(1, 1, 1)
    c.setFont("Helvetica-Bold", 22)
    c.drawString(60, h - 58, title)
    c.setFillColorRGB(0, 0, 0)
    c.setFont("Helvetica", 13)
    y = h - 140
    for line in lines:
        if line.startswith("#"):
            c.setFont("Helvetica-Bold", 15); y -= 6
            c.drawString(60, y, line[1:].strip()); c.setFont("Helvetica", 13); y -= 26
        else:
            c.drawString(60, y, line); y -= 21
    c.setFont("Helvetica-Oblique", 9)
    c.drawString(60, 40, f"Enhanced Task Manager - {filename}")
    c.showPage()
    c.save()


make_pdf("01_project_overview.pdf", "1. Project Overview", [
    "Project: Enhanced Task Manager",
    "Goal: a task manager with priorities, due dates, search and reports.",
    "",
    "# Planning",
    "- Requirements gathering and user stories",
    "- Database schema design (tasks, users, tags)",
    "- Modular design: storage / logic / reporting",
    "",
    "# Design",
    "- Task: id, title, priority (High/Medium/Low), due date, status",
    "- Config-driven automation for documents (this lab)",
])
make_pdf("02_task_list.pdf", "2. Task List", [
    "T-01  Design database schema              High     Done",
    "T-02  Create / edit / delete tasks        High     Done",
    "T-03  Priority & due-date filter          Medium   Done",
    "T-04  Search tasks by keyword             Medium   Done",
    "T-05  Weekly PDF/Word report              Medium   In progress",
    "T-06  Email reminders                     Low      To do",
])
make_pdf("03_progress_summary.pdf", "3. Progress Summary", [
    "Completed: 4 of 6 tasks (67%)",
    "",
    "# Next steps",
    "- Finish weekly report generation",
    "- Add email reminders for due tasks",
    "- Write unit tests for the task module",
])

# Word template with {{placeholders}}
d = docx.Document()
d.add_heading("{{report_title}}", level=0)
p = d.add_paragraph()
p.add_run("Student: ").bold = True
p.add_run("{{student_name}}")
p = d.add_paragraph()
p.add_run("Student ID: ").bold = True
p.add_run("{{student_id}}")
p = d.add_paragraph()
p.add_run("Date: ").bold = True
p.add_run("{{date}}")
d.add_heading("Project Summary", level=1)
d.add_paragraph("Project: {{project_name}}")
d.add_paragraph("Progress: {{completed_tasks}} of {{total_tasks}} tasks completed.")
d.add_paragraph("{{summary}}")
for para in d.paragraphs:
    for run in para.runs:
        run.font.size = Pt(12) if run.font.size is None else run.font.size
d.save(os.path.join(TEMPLATES, "report_template.docx"))

print("Created 3 sample PDFs in documents/ and templates/report_template.docx")
