# src/word_tasks.py
import datetime
import os
import re

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from .utils import setup_logging, ensure_directory_exists

logger = setup_logging(__name__)

PLACEHOLDER = re.compile(r"\{\{\s*(\w+)\s*\}\}")


class WordTasks:
    """Word operations: generating a report from a .docx template with {{placeholders}}."""

    @staticmethod
    def _replace_in_paragraph(paragraph, context):
        """Replaces {{key}} in a paragraph, even if Word split the placeholder over several runs."""
        full_text = "".join(run.text for run in paragraph.runs)
        if "{{" not in full_text:
            return 0
        new_text, count = PLACEHOLDER.subn(lambda m: str(context.get(m.group(1), m.group(0))), full_text)
        if count:
            paragraph.runs[0].text = new_text  # keep the formatting of the first run
            for run in paragraph.runs[1:]:
                run.text = ""
        return count

    def generate_report(self, template, output_file, context, table=None):
        """Fills the template, appends an optional data table, and saves the report."""
        if not os.path.exists(template):
            raise FileNotFoundError(f"Template not found: {template}")
        context = dict(context)
        context.setdefault("date", datetime.date.today().isoformat())
        document = docx.Document(template)

        replaced = 0
        for p in document.paragraphs:
            replaced += self._replace_in_paragraph(p, context)
        for t in document.tables:
            for row in t.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        replaced += self._replace_in_paragraph(p, context)
        logger.info(f"Filled {replaced} placeholder(s) in template '{os.path.basename(template)}'.")

        if table:
            document.add_heading(table.get("title", "Data"), level=2)
            rows = [table["headers"]] + table["rows"]
            t = document.add_table(rows=len(rows), cols=len(table["headers"]))
            t.style = "Light Grid Accent 1"
            for r, row in enumerate(rows):
                for c, value in enumerate(row):
                    cell = t.cell(r, c)
                    cell.text = str(value)
                    if r == 0:
                        cell.paragraphs[0].runs[0].bold = True
            logger.info(f"Added table '{table.get('title')}' with {len(table['rows'])} row(s).")

        footer = document.add_paragraph(f"Generated automatically by the Document Automation Agent on {context['date']}")
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.runs[0].italic = True
        footer.runs[0].font.size = Pt(9)
        footer.runs[0].font.color.rgb = RGBColor(0x80, 0x80, 0x80)

        ensure_directory_exists(os.path.dirname(output_file))
        document.save(output_file)
        logger.info(f"Report saved: {output_file}")
        return replaced
